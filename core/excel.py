"""Отчёт за период в Excel: те же цифры, что на сайте, по листам."""
import io

from django.utils import timezone
from django.utils.translation import gettext as _
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import models as m
from . import reports

MONEY_FMT = "#,##0"
QTY_FMT = "#,##0.###"
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")


def _sheet(wb, title, header, rows, money_cols=(), qty_cols=()):
    ws = wb.create_sheet(title)
    ws.append(header)
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.fill = HEAD_FILL
    for row in rows:
        ws.append([float(v) if hasattr(v, "quantize") else v for v in row])
    widths = [len(str(h)) for h in header]
    for row in ws.iter_rows(min_row=2):
        for i, cell in enumerate(row):
            if i in money_cols:
                cell.number_format = MONEY_FMT
            elif i in qty_cols:
                cell.number_format = QTY_FMT
            widths[i] = max(widths[i], len(str(cell.value or "")) + 2)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = min(max(w + 2, 10), 45)
    ws.freeze_panes = "A2"
    return ws


def build(period):
    start = reports.period_start(period)
    data = reports.build(period)
    cash = reports.money(period)
    staff = reports.staff(period)["rows"]
    today = timezone.localdate()

    wb = Workbook()
    wb.remove(wb.active)
    since = f"{start:%d.%m.%Y}" if start else "—"
    _sheet(wb, _("Итоги"), [_("Показатель"), _("Сумма")], [
        (_("Период"), f"{reports.PERIODS[period]}: {since} – {today:%d.%m.%Y}"),
        (_("Продажи"), data["revenue"]),
        (_("Возвраты"), data["returned"]),
        (_("Себестоимость проданного"), data["cogs"]),
        (_("Прочие расходы"), data["expenses"]),
        (_("Прибыль"), data["profit"]),
        (_("Пришло денег"), cash["cash_in"]),
        (_("Из них картой"), cash["card_in"]),
        (_("Закуплено сырья"), cash["purchased"]),
        (_("Нам должны"), cash["receivable"]),
        (_("Мы должны"), cash["payable"]),
        (_("Продано, шт"), data["sold_qty"]),
        (_("Возвращено, шт"), data["returned_qty"]),
    ], money_cols=(1,))

    _sheet(wb, _("Магазины"),
           [_("Магазин"), _("Район"), _("Продажи"), _("Шт"), _("Возвраты, шт"), _("Возврат, %"), _("Долг")],
           [(s["shop"].name, str(s["shop"].district or ""), s["sales"], s["qty"], s["returns"], s["return_pct"],
             s["debt"]) for s in data["shops"]],
           money_cols=(2, 6), qty_cols=(3, 4))

    _sheet(wb, _("Продукты"),
           [_("Продукт"), _("Продажи"), _("Шт"), _("Возвраты, шт"), _("Цена"), _("Себестоимость"),
            _("Прибыль с 1 шт"), _("На складе")],
           [(p["product"].name, p["sales"], p["qty"], p["returns"], p["product"].price, p["unit_cost"],
             p["unit_profit"], p["stock"]) for p in data["products"]],
           money_cols=(1, 4, 5, 6), qty_cols=(2, 3, 7))

    _sheet(wb, _("Сотрудники"),
           [_("Сотрудник"), _("Роль"), _("Продажи"), _("Шт"), _("Магазинов"), _("Собрал денег"), _("Расходы"),
            _("Произвёл, шт"), _("Возвраты, шт"), _("Дней"), _("Оклад"), _("Выдано за месяц")],
           [(r["worker"].name, r["worker"].get_role_display(), r["sales"], r["qty"], r["shops"], r["cash"],
             r["expenses"], r["produced"], r["returns"], r["days"], r["salary"], r["salary_paid"])
            for r in staff],
           money_cols=(2, 5, 6, 10, 11), qty_cols=(3, 7, 8))

    items = reports._since(m.PurchaseItem.objects.select_related("purchase__supplier", "material__unit"),
                           "purchase__date", start).order_by("purchase__date", "id")
    _sheet(wb, _("Закупки"),
           [_("Дата"), _("Поставщик"), _("Сырьё"), _("Количество"), _("Ед."), _("Цена"), _("Сумма")],
           [(f"{it.purchase.date:%d.%m.%Y}", it.purchase.supplier.name, it.material.name, it.qty,
             str(it.material.unit), it.price, it.qty * it.price) for it in items],
           money_cols=(5, 6), qty_cols=(3,))

    expenses = reports._since(m.Expense.objects.select_related("worker"), "date", start).order_by("date", "id")
    _sheet(wb, _("Расходы"),
           [_("Дата"), _("Статья"), _("Сумма"), _("Сотрудник"), _("Заметка")],
           [(f"{e.date:%d.%m.%Y}", e.get_category_display(), e.amount, e.worker.name if e.worker else "", e.note)
            for e in expenses],
           money_cols=(2,))

    _sheet(wb, _("Долги"), [_("Кто"), _("Тип"), _("Сумма")],
           [(s.name, _("Магазин должен нам"), d) for s, d in cash["debtors"]]
           + [(s.name, _("Мы должны поставщику"), d) for s, d in cash["supplier_debts"]],
           money_cols=(2,))

    out = io.BytesIO()
    wb.save(out)
    name = f"dastyor-{period}-{today:%Y-%m-%d}.xlsx"
    return name, out.getvalue()
